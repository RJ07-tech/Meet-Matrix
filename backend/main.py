import os
import uuid
import io
import openpyxl
import hmac
from fastapi.responses import Response
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
from typing import Optional, Dict, Any, List, Set
from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, ConfigDict
from livekit import api
from livekit.api import AccessToken, VideoGrants

app = FastAPI(title="MeetMatrix Backend API")

# Indian Standard Time (IST = UTC + 5:30)
IST = timezone(timedelta(hours=5, minutes=30))

def get_ist_now_dt() -> datetime:
    return datetime.now(IST)

def get_ist_now_str() -> str:
    return get_ist_now_dt().strftime("%Y-%m-%d %I:%M:%S %p")

_cors_raw = os.getenv("FRONTEND_ORIGINS", "*").strip()
if _cors_raw == "*":
    _cors_origins = ["*"]
    _cors_credentials = False
else:
    _cors_origins = [o.strip() for o in _cors_raw.split(",") if o.strip()]
    _cors_credentials = True

origins = [
    "https://meet-matrix.vercel.app",
    "http://localhost:5173",
    "http://localhost:3000",
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,  # Explicit domains allowed with credentials
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

LIVEKIT_URL = os.getenv("LIVEKIT_URL", "https://meetmatrix-xxxxxx.livekit.cloud")
LIVEKIT_API_KEY = os.getenv("LIVEKIT_API_KEY", "devkey")
LIVEKIT_API_SECRET = os.getenv("LIVEKIT_API_SECRET", "secret")
TOKEN_TTL_HOURS = int(os.getenv("LIVEKIT_TOKEN_TTL_HOURS", "6"))

room_settings_db: Dict[str, Dict[str, Any]] = {}
scheduled_meetings_db: Dict[str, Dict[str, Any]] = {}
waiting_room_db: Dict[str, List[Dict[str, Any]]] = {}
banned_participants_db: Dict[str, Dict[str, Set[str]]] = {}
host_secrets_db: Dict[str, str] = {}
attendance_db: Dict[str, List[Dict[str, Any]]] = {}


class TokenRequest(BaseModel):
    room_name: str
    participant_name: str
    is_host: bool = False
    role: Optional[str] = "participant"
    host_secret: Optional[str] = None
    participant_identity: Optional[str] = None


class RoomConfig(BaseModel):
    room_id: Optional[str] = None
    waiting_mode: str = "direct"
    chat_locked: bool = False
    chat_host_only: bool = False
    mic_locked: bool = False
    allow_participant_screenshare: bool = False
    allow_cohost_whiteboard: bool = True
    allow_direct_chat: bool = False
    mute_on_entry: bool = False
    camera_off_on_entry: bool = False
    allow_whiteboard: bool = False
    allow_reactions: bool = False
    auto_download_Excel: bool = True


class UpdateRoomSettingsRequest(BaseModel):
    model_config = ConfigDict(extra="ignore")
    room_name: str
    waiting_mode: Optional[str] = None
    chat_locked: Optional[bool] = None
    chat_host_only: Optional[bool] = None
    mic_locked: Optional[bool] = None
    allow_participant_screenshare: Optional[bool] = None
    allow_cohost_whiteboard: Optional[bool] = None
    allow_direct_chat: Optional[bool] = None
    mute_on_entry: Optional[bool] = None
    camera_off_on_entry: Optional[bool] = None
    allow_whiteboard: Optional[bool] = None
    allow_reactions: Optional[bool] = None
    auto_download_Excel: Optional[bool] = None


class AdmitRequest(BaseModel):
    room_name: str
    participant_id: str
    action: str


class KickRequest(BaseModel):
    room_name: str
    participant_identity: str
    participant_name: Optional[str] = None


class TerminateRequest(BaseModel):
    room_name: str


class AttendanceUpdateRequest(BaseModel):
    room_name: str
    participant_name: str
    participant_identity: Optional[str] = None
    action: str  # "join", "leave", "hold_start", "hold_end"


def new_room_id() -> str:
    return f"mm-{uuid.uuid4().hex[:8]}-{uuid.uuid4().hex[:8]}"


def default_settings(room_id: str) -> Dict[str, Any]:
    return RoomConfig(room_id=room_id).model_dump()


def settings_from_payload(room_id: str, data: Dict[str, Any]) -> Dict[str, Any]:
    base = default_settings(room_id)
    for key in RoomConfig.model_fields:
        if key == "room_id":
            continue
        if key in data and data[key] is not None:
            base[key] = data[key]
    base["room_id"] = room_id
    return base


def public_settings(room_name: str) -> Dict[str, Any]:
    return dict(room_settings_db.get(room_name) or default_settings(room_name))


def public_meeting(meeting: Dict[str, Any]) -> Dict[str, Any]:
    return {k: v for k, v in meeting.items() if k != "host_secret"}


def ensure_bans(room_name: str) -> Dict[str, Set[str]]:
    entry = banned_participants_db.get(room_name)
    if entry is None:
        banned_participants_db[room_name] = {"identities": set(), "names": set()}
    elif isinstance(entry, set):
        banned_participants_db[room_name] = {"identities": set(), "names": set(entry)}
    return banned_participants_db[room_name]


def normalize_name(name: Optional[str]) -> str:
    return (name or "").replace("(Host)", "").replace("(Co-Host)", "").strip().lower()


def names_match(stored: Optional[str], clean_name: str) -> bool:
    return normalize_name(stored) == clean_name


def verified_host(room_name: str, host_secret: Optional[str]) -> bool:
    stored = host_secrets_db.get(room_name)
    if not stored or not host_secret or len(stored) != len(host_secret):
        return False
    return hmac.compare_digest(stored, host_secret)


def find_attendance(records, identity=None, name=None):
    if identity:
        rec = next((r for r in records if r.get("identity") == identity), None)
        if rec:
            return rec
    if name:
        rec = next((r for r in records if names_match(r.get("name"), normalize_name(name))), None)
        if rec:
            return rec
    return None


def mark_attendance_join(room_name: str, name: str, identity: str, is_host: bool = False) -> None:
    if room_name not in attendance_db:
        attendance_db[room_name] = []

    records = attendance_db[room_name]
    now_dt = get_ist_now_dt()
    now_str = now_dt.strftime("%Y-%m-%d %I:%M:%S %p")
    rec = find_attendance(records, identity=identity, name=name)

    initial_role = "Host" if is_host else "Attendee"

    if rec:
        rec["leave_time"] = "Active"
        rec["identity"] = identity or rec.get("identity")

        # Track initial name and any renames
        if name and name != rec.get("name"):
            rec.setdefault("name_history", [rec.get("name", name)])
            if name not in rec["name_history"]:
                rec["name_history"].append(name)
            rec["name"] = name

        if not rec.get("join_dt"):
            rec["join_dt"] = now_dt

        if is_host:
            rec["role"] = "Host"
            rec["is_host"] = True
            if "Host" not in rec.setdefault("roles_held", []):
                rec["roles_held"].append("Host")

        return

    records.append({
        "name": name,
        "identity": identity,
        "is_host": is_host,
        "role": initial_role,
        "roles_held": [initial_role],
        "name_history": [name],  # Track all names used by this participant
        "join_time": now_str,
        "join_dt": now_dt,
        "leave_time": "Active",
        "current_hold_start": None,
        "hold_count": 0,
        "total_hold_seconds": 0,
        "total_session_seconds": 0,
        "active_seconds": 0,
        "hold_logs": [],
    })


def is_room_live(room_name: str) -> bool:
    records = attendance_db.get(room_name) or []
    return any(r.get("leave_time") == "Active" for r in records)


@app.post("/api/create-room")
async def create_room(cfg: RoomConfig):
    rid = cfg.room_id if cfg.room_id else new_room_id()
    room_settings_db[rid] = settings_from_payload(rid, cfg.model_dump())
    host_secrets_db[rid] = uuid.uuid4().hex
    ensure_bans(rid)
    if rid not in attendance_db:
        attendance_db[rid] = []
    return {
        "status": "success",
        "room_id": rid,
        "host_secret": host_secrets_db[rid],
        "settings": public_settings(rid),
    }


@app.get("/api/room-settings/{room_name}")
async def get_room_settings(room_name: str):
    return public_settings(room_name)


@app.post("/api/update-room-settings")
async def update_room_settings(req: UpdateRoomSettingsRequest):
    rname = req.room_name
    current = room_settings_db.setdefault(rname, default_settings(rname))
    updates = req.model_dump(exclude={"room_name"}, exclude_unset=True)
    current.update(updates)
    current["room_id"] = rname
    room_settings_db[rname] = current
    return {"status": "success", "settings": public_settings(rname)}


@app.post("/api/get-token")
async def get_token(req: TokenRequest):
    cfg = public_settings(req.room_name)
    waiting_mode = cfg.get("waiting_mode", "direct")
    clean_name = normalize_name(req.participant_name)
    bans = ensure_bans(req.room_name)
    is_host = verified_host(req.room_name, req.host_secret)

    identity_banned = bool(req.participant_identity) and req.participant_identity in bans["identities"]
    name_banned = clean_name in bans["names"] if clean_name else False
    is_banned = identity_banned or name_banned

    must_wait = (waiting_mode == "strict" or is_banned) and not is_host

    if must_wait:
        if req.room_name not in waiting_room_db:
            waiting_room_db[req.room_name] = []

        current_list = waiting_room_db[req.room_name]
        admitted_index = next(
            (i for i, p in enumerate(current_list) if names_match(p.get("name"), clean_name) and p.get("status") == "admitted"),
            None
        )

        # If not admitted yet by host, put in queue and block entry
        if admitted_index is None:
            existing = next(
                (p for p in current_list if names_match(p.get("name"), clean_name)),
                None
            )
            if not existing:
                new_pid = str(uuid.uuid4())
                current_list.append({
                    "participant_id": new_pid,
                    "name": req.participant_name,
                    "status": "waiting"
                })
            else:
                new_pid = existing["participant_id"]
                existing["status"] = "waiting"

            return {
                "status": "waiting",
                "participant_id": new_pid,
                "message": "Waiting for host admission"
            }
        else:
            # Participant was admitted by host; consume and pop the ticket so rejoins cannot bypass
            current_list.pop(admitted_index)

    try:
        grants = VideoGrants(
            room_join=True,
            room=req.room_name,
            can_publish=True,
            can_subscribe=True,
            can_publish_data=True
        )

        pid = req.participant_identity if req.participant_identity and not identity_banned else f"{req.participant_name}_{uuid.uuid4().hex[:8]}"
        token = AccessToken(LIVEKIT_API_KEY, LIVEKIT_API_SECRET) \
            .with_identity(pid) \
            .with_name(req.participant_name) \
            .with_grants(grants) \
            .with_ttl(timedelta(hours=TOKEN_TTL_HOURS))

        jwt_token = token.to_jwt()
        mark_attendance_join(req.room_name, req.participant_name, pid, is_host=is_host)

        return {
            "status": "success",
            "token": jwt_token,
            "server_url": LIVEKIT_URL,
            "identity": pid,
            "is_host": is_host,
            "mute_on_entry": cfg.get("mute_on_entry", False),
            "camera_off_on_entry": cfg.get("camera_off_on_entry", False),
            "mic_locked": cfg.get("mic_locked", False)
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/attendance/update")
async def update_attendance(req: AttendanceUpdateRequest):
    if req.room_name not in attendance_db:
        attendance_db[req.room_name] = []

    records = attendance_db[req.room_name]
    rec = find_attendance(records, identity=req.participant_identity, name=req.participant_name)
    now_dt = get_ist_now_dt()
    now_str = now_dt.strftime("%Y-%m-%d %I:%M:%S %p")

    if req.action == "join":
        mark_attendance_join(
            req.room_name,
            req.participant_name,
            req.participant_identity or "Unknown",
            is_host=False,
            )
        return {"status": "success"}

    if not rec:
        mark_attendance_join(
            req.room_name,
            req.participant_name,
            req.participant_identity or "Unknown",
            is_host=False,
            )
        rec = find_attendance(records, identity=req.participant_identity, name=req.participant_name)

    if rec:
        rec.setdefault("roles_held", [rec.get("role", "Attendee")])
        rec.setdefault("name_history", [rec.get("name", req.participant_name)])

        # If a participant submits a name change or changes name during an event
        if req.participant_name and req.participant_name != rec.get("name"):
            if req.participant_name not in rec["name_history"]:
                rec["name_history"].append(req.participant_name)
            rec["name"] = req.participant_name

        if req.action == "name_change":
            # Explicit name change event
            return {"status": "success"}

    if rec:
        # Guarantee roles_held exists
        if "roles_held" not in rec or not isinstance(rec["roles_held"], list):
            rec["roles_held"] = [rec.get("role", "Attendee")]

        if req.action == "role_cohost":
            if rec.get("role") != "Host":
                rec["role"] = "Co-Host"
                if "Co-Host" not in rec["roles_held"]:
                    rec["roles_held"].append("Co-Host")
                # Clear active hold if they were promoted while away
                rec["current_hold_start"] = None

        elif req.action == "role_attendee":
            if rec.get("role") != "Host":
                rec["role"] = "Attendee"
                # Keep "Co-Host" in rec["roles_held"]! Do NOT remove it.
        elif req.action == "hold_start":
            if not rec.get("current_hold_start") and rec.get("role") == "Attendee":
                rec["current_hold_start"] = now_dt
                rec["hold_count"] = rec.get("hold_count", 0) + 1

        elif req.action == "hold_end":
            if rec.get("current_hold_start"):
                start_dt = rec["current_hold_start"]
                dur = max(0, int((now_dt - start_dt).total_seconds()))
                rec.setdefault("hold_logs", []).append(
                    f"{start_dt.strftime('%I:%M:%S %p')} - {now_dt.strftime('%I:%M:%S %p')} ({dur}s)"
                )
                rec["total_hold_seconds"] = rec.get("total_hold_seconds", 0) + dur
                rec["current_hold_start"] = None

        elif req.action == "leave":
            rec["leave_time"] = now_str
            if rec.get("current_hold_start"):
                start_dt = rec["current_hold_start"]
                dur = max(0, int((now_dt - start_dt).total_seconds()))
                rec.setdefault("hold_logs", []).append(
                    f"{start_dt.strftime('%I:%M:%S %p')} - {now_dt.strftime('%I:%M:%S %p')} ({dur}s)"
                )
                rec["total_hold_seconds"] = rec.get("total_hold_seconds", 0) + dur
                rec["current_hold_start"] = None

            join_dt = rec.get("join_dt", now_dt)
            total_sec = max(1, int((now_dt - join_dt).total_seconds()))
            hold_sec = rec.get("total_hold_seconds", 0)
            active_sec = max(0, total_sec - hold_sec)

            rec["total_session_seconds"] = total_sec
            rec["active_seconds"] = active_sec
            rec["active_percentage"] = round((active_sec / total_sec) * 100, 1)

    return {"status": "success"}

def format_duration(seconds: int) -> str:
    m, s = divmod(int(seconds), 60)
    h, m = divmod(m, 60)
    if h > 0:
        return f"{h}h {m}m {s}s"
    return f"{m}m {s}s"

@app.get("/api/attendance/export/{room_name}")
async def export_attendance_excel(room_name: str):
    records = attendance_db.get(room_name, [])
    now_dt = get_ist_now_dt()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    ws.views.sheetView[0].showGridLines = True

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="0F172A")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "Participant Name",
        "Role During Meeting",
        "Join Time",
        "Leave Time",
        "Total Meeting Duration",
        "Away / Hold Time",
        "Active Attention Time",
        "Name Change History",
        "Away Instances",
        "Detailed Away Intervals"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    for row_idx, rec in enumerate(records, start=2):
        join_dt = rec.get("join_dt", now_dt)
        hold_sec = rec.get("total_hold_seconds", 0)

        if rec.get("current_hold_start"):
            hold_sec += max(0, int((now_dt - rec["current_hold_start"]).total_seconds()))

        total_sec = rec.get("total_session_seconds") or max(1, int((now_dt - join_dt).total_seconds()))
        active_sec = max(0, total_sec - hold_sec)

        # Build clean role description
        current_role = rec.get("role", "Attendee")
        roles_held = rec.get("roles_held", [])
        if current_role == "Host" or "Host" in roles_held:
            display_role = "Host"
        elif "Co-Host" in roles_held:
            display_role = "Co-Host" if current_role == "Co-Host" else "Attendee (Was Co-Host)"
        else:
            display_role = current_role

        # Build Name Change description
        name_history = rec.get("name_history", [rec.get("name", "Unknown")])
        if len(name_history) > 1:
            name_change_str = f"Yes ({' -> '.join(name_history)})"
        else:
            name_change_str = "No"

        hold_logs_str = "\n".join(rec.get("hold_logs", [])) if rec.get("hold_logs") else "None"

        row_data = [
            rec.get("name", "Unknown"),
            display_role,
            rec.get("join_time", "-"),
            rec.get("leave_time", "-"),
            format_duration(total_sec),
            format_duration(hold_sec),
            format_duration(active_sec),
            name_change_str,
            rec.get("hold_count", 0),
            hold_logs_str
        ]

        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 8, 10]:
                cell.alignment = left_align
            else:
                cell.alignment = center_align

            if col_idx in [8, 10]:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws.row_dimensions[row_idx].height = 24 if not rec.get("hold_logs") else 20 * max(1, len(rec.get("hold_logs", [])))

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    return Response(
        content=file_stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=attendance_{room_name}.xlsx"}
    )


@app.get("/api/waiting-list/{room_name}")
async def get_waiting_list(room_name: str):
    pending = [
        p for p in waiting_room_db.get(room_name, [])
        if p.get("status") == "waiting"
    ]
    return {"waiting": pending}


@app.post("/api/admit-participant")
async def admit_participant(req: AdmitRequest):
    r_list = waiting_room_db.get(req.room_name, [])
    for p in r_list:
        if p.get("participant_id") == req.participant_id:
            p["status"] = "admitted" if req.action == "admit" else "rejected"
            if req.action == "admit":
                bans = ensure_bans(req.room_name)
                bans["names"].discard(normalize_name(p.get("name")))
                if p.get("identity"):
                    bans["identities"].discard(p["identity"])
            return {"status": "success", "action": req.action}
    raise HTTPException(status_code=404, detail="Participant not found")


@app.get("/api/check-admission/{room_name}/{participant_id}")
async def check_admission(room_name: str, participant_id: str):
    r_list = waiting_room_db.get(room_name, [])
    for p in r_list:
        if p.get("participant_id") == participant_id:
            return {"status": p.get("status", "waiting")}
    return {"status": "not_found"}


@app.post("/api/schedule-meeting")
async def schedule_meeting(data: Dict[str, Any]):
    rid = data.get("room_id") or new_room_id()
    data["room_id"] = rid
    secret = uuid.uuid4().hex
    host_secrets_db[rid] = secret
    data["host_secret"] = secret
    scheduled_meetings_db[rid] = data
    room_settings_db[rid] = settings_from_payload(rid, data)
    return {"status": "success", "room_id": rid, "host_secret": secret, "meeting": public_meeting(data)}


@app.get("/api/scheduled-meetings")
async def list_scheduled_meetings():
    return {"meetings": [public_meeting(m) for m in scheduled_meetings_db.values()]}


@app.delete("/api/scheduled-meetings/{room_id}")
async def delete_scheduled_meeting(room_id: str):
    if room_id not in scheduled_meetings_db:
        return {"status": "not_found"}
    del scheduled_meetings_db[room_id]
    if not is_room_live(room_id):
        room_settings_db.pop(room_id, None)
        host_secrets_db.pop(room_id, None)
        waiting_room_db.pop(room_id, None)
        banned_participants_db.pop(room_id, None)
    return {"status": "deleted"}


@app.post("/api/kick-participant")
async def kick_participant(req: KickRequest):
    bans = ensure_bans(req.room_name)
    if req.participant_identity:
        bans["identities"].add(req.participant_identity)
    clean_target = normalize_name(req.participant_name)
    if clean_target:
        bans["names"].add(clean_target)

    if req.room_name in waiting_room_db:
        for p in waiting_room_db[req.room_name]:
            if names_match(p.get("name"), clean_target) or p.get("identity") == req.participant_identity:
                p["status"] = "rejected"

    lk_api = api.LiveKitAPI(LIVEKIT_URL, LIVEKIT_API_KEY, LIVEKIT_API_SECRET)
    try:
        await lk_api.room.remove_participant(
            api.RoomParticipantIdentity(
                room=req.room_name,
                identity=req.participant_identity
            )
        )
        return {"status": "success"}
    except Exception as e:
        return {"status": "acknowledged", "detail": str(e)}
    finally:
        await lk_api.aclose()


@app.post("/api/terminate-room")
async def terminate_room(req: TerminateRequest):
    # Mark leave times before terminating room
    if req.room_name in attendance_db:
        now_dt = get_ist_now_dt()
        now_str = now_dt.strftime("%Y-%m-%d %I:%M:%S %p")
        for r in attendance_db[req.room_name]:
            if r.get("leave_time") == "Active":
                r["leave_time"] = f"{now_str} (Meeting Ended)"
            if r.get("current_hold_start"):
                s_dt = r["current_hold_start"]
                dur = int((now_dt - s_dt).total_seconds())
                r.setdefault("hold_logs", []).append(
                    f"{s_dt.strftime('%I:%M:%S %p')} to {now_dt.strftime('%I:%M:%S %p')} ({dur}s)"
                )
                r["current_hold_start"] = None

    lk_api = api.LiveKitAPI(LIVEKIT_URL, LIVEKIT_API_KEY, LIVEKIT_API_SECRET)
    try:
        await lk_api.room.delete_room(api.DeleteRoomRequest(room=req.room_name))
    except Exception:
        pass
    finally:
        await lk_api.aclose()

    room_settings_db.pop(req.room_name, None)
    waiting_room_db.pop(req.room_name, None)
    banned_participants_db.pop(req.room_name, None)
    host_secrets_db.pop(req.room_name, None)
    scheduled_meetings_db.pop(req.room_name, None)
    return {"status": "terminated"}


@app.get("/health")
async def health():
    return {"status": "healthy"}